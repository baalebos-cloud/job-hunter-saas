import os
os.environ['DATABASE_URL']='sqlite://'
os.environ['SECRET_KEY']='test-key-only-not-for-production'
import io
import json
from types import SimpleNamespace
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from backend.app.main import app
from backend.app.database import Base, get_db
from backend.app.models.user import User
from backend.app.models.job import Job
from backend.app.models.resume import Resume
from backend.app.models.application import Application
from backend.app.models.subscription import Subscription
from backend.app.models.referral import Referral
from backend.app.services.auth_service import create_access_token
from backend.app.utils.resume_quality import merge_rewrite, score_requirements
from backend.app.utils.pdf_generator import generate_optimized_resume
from backend.app.utils.ats_engine import extract_text

@pytest.fixture
def db():
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session=sessionmaker(bind=engine)()
    users=[User(email='jayeolaoluwadamilare@gmail.com',full_name='Owner',hashed_password='x',is_admin=True,is_verified=True),
           User(email='candidate@example.com',full_name='Candidate',hashed_password='x',is_verified=True),
           User(email='hr@example.com',full_name='HR',hashed_password='x',is_hr=True,is_verified=True,hr_approved=True),
           User(email='pending@example.com',hashed_password='x',is_hr=True,is_verified=True,hr_approved=False)]
    session.add_all(users);session.commit()
    app.dependency_overrides[get_db]=lambda:session
    yield session
    app.dependency_overrides.clear();session.close();engine.dispose()

@pytest.fixture
def client(db): return TestClient(app)
def headers(email='candidate@example.com'): return {'Authorization':'Bearer '+create_access_token({'sub':email})}

def test_owner_and_hr_boundaries(client,db):
    assert client.get('/api/v1/admin/users').status_code==401
    assert client.get('/api/v1/admin/users',headers=headers()).status_code==403
    assert client.get('/api/v1/admin/users',headers=headers('jayeolaoluwadamilare@gmail.com')).status_code==200
    assert client.post('/api/v1/hr-auth/approve/4').status_code==401
    assert client.get('/api/v1/hr-auth/pending').status_code==401
    assert client.get('/api/v1/hr/dashboard',headers=headers('pending@example.com')).status_code==403
    job=Job(title='Role',company='Employer',hr_user_id=1,posted_by_hr=True,is_active=True);db.add(job);db.commit()
    application=Application(user_id=2,job_id=job.id,status='submitted_internal');db.add(application);db.commit()
    assert client.get(f'/api/v1/hr/jobs/{job.id}/applications',headers=headers('hr@example.com')).status_code==404
    assert client.patch(f'/api/v1/hr/applications/{application.id}/status?new_status=offer',headers=headers('hr@example.com')).status_code==404
    assert client.patch(f'/api/v1/admin/users/2/make-admin',headers=headers('jayeolaoluwadamilare@gmail.com')).status_code==403


def test_signup_cannot_self_assign_hr(client,db,monkeypatch):
    monkeypatch.setattr('backend.app.routes.auth.send_verification_email',lambda **kw:None)
    res=client.post('/api/v1/auth/signup',json={'full_name':'Test','email':'new@example.com','password':'example-password','is_hr':True})
    assert res.status_code==201
    assert not db.query(User).filter_by(email='new@example.com').first().is_hr


def test_resume_privacy_and_pdf_preview(client,db):
    db.add(Resume(owner_id=2,filename='optimized_demo.pdf',content=b'%PDF-example',analysis_data={}));db.commit()
    for endpoint in ('preview','download','status'):
        assert client.get(f'/api/v1/resume/{endpoint}/demo').status_code==401
        assert client.get(f'/api/v1/resume/{endpoint}/demo',headers=headers('hr@example.com')).status_code==404
    res=client.get('/api/v1/resume/preview/demo',headers=headers())
    assert res.status_code==200 and res.headers['content-type']=='application/pdf'
    assert res.headers['content-disposition'].startswith('inline')


def test_external_handoff_is_not_submission(client,db):
    job=Job(title='DevOps',company='Example',description='Python and AWS experience required.',url='https://example.com/apply',is_active=True)
    db.add(job);db.commit()
    res=client.post(f'/api/v1/jobs/{job.id}/apply',json={},headers=headers())
    assert res.status_code==200 and res.json()['status']=='external_started'
    record=db.query(Application).first()
    assert record.ats_score is None and record.job_snapshot['description']==job.description
    res=client.patch(f'/api/v1/jobs/{job.id}/application-status?new_status=submitted_external',headers=headers())
    assert res.status_code==200
    assert client.post(f'/api/v1/jobs/{job.id}/apply',json={},headers=headers()).json()['status']=='submitted_external'


def test_internal_requires_correct_resume_and_preserves_attachment(client,db):
    job=Job(title='Python',company='Employer',description='Python application development required.',hr_user_id=3,posted_by_hr=True,is_active=True)
    db.add(job);db.commit()
    db.add(Resume(owner_id=2,filename='optimized_exact.pdf',content=b'%PDF-example',analysis_data={'job_description':job.description,'overall_score':50}));db.commit()
    assert client.post(f'/api/v1/jobs/{job.id}/apply',json={},headers=headers()).status_code==422
    res=client.post(f'/api/v1/jobs/{job.id}/apply',json={'resume_id':'exact'},headers=headers())
    assert res.status_code==200 and res.json()['status']=='submitted_internal'
    assert db.query(Application).first().resume_id is not None
    application_id = db.query(Application).first().id
    assert client.get(f'/api/v1/hr/applications/{application_id}/resume', headers=headers('hr@example.com')).status_code == 200
    assert client.get(f'/api/v1/hr/applications/{application_id}/resume', headers=headers('pending@example.com')).status_code == 403
    assert client.get(f'/api/v1/hr/applications/{application_id}/resume', headers=headers()).status_code == 403


def test_successful_upload_scores_exported_pdf_and_charges_once(client, db, monkeypatch):
    import backend.app.utils.ats_engine as engine
    base = original()
    monkeypatch.setattr(engine, 'extract_resume_data', lambda text: base)
    monkeypatch.setattr(engine, 'extract_requirements', lambda jd: [{'phrase': 'Python'}, {'phrase': 'Terraform'}])
    monkeypatch.setattr(engine, 'rewrite_resume_for_job', lambda **kw: {
        'optimized_summary': base['summary'], 'optimized_experience': base['experience'], 'optimized_skills': base['skills']})
    input_pdf = generate_optimized_resume('input.pdf', resume_data=base).getvalue()
    response = client.post('/api/v1/resume/upload', headers=headers(),
        files={'file': ('input.pdf', input_pdf, 'application/pdf')},
        data={'job_title': 'Engineer', 'job_description': 'Python and Terraform skills required for this engineering role.'})
    assert response.status_code == 200, response.text
    result = response.json()['result']
    assert result['overall_score'] == 50
    record = db.query(Resume).first()
    assert json.loads(record.parsed_data)['projects'] == base['projects']
    assert db.query(Subscription).first().scans_this_month == 1
    assert client.get('/api/v1/resume/preview/' + response.json()['task_id'], headers=headers()).content == record.content


def test_failed_upload_does_not_consume_quota(client, db, monkeypatch, caplog):
    def unavailable(text):
        raise RuntimeError('Provider unavailable')
    monkeypatch.setattr('backend.app.utils.ats_engine.extract_resume_data', unavailable)
    input_pdf = generate_optimized_resume('input.pdf', resume_data=original()).getvalue()
    response = client.post('/api/v1/resume/upload', headers=headers(),
        files={'file': ('input.pdf', input_pdf, 'application/pdf')},
        data={'job_title': 'Engineer', 'job_description': 'Python experience required for this engineering role.'})
    assert response.status_code == 503
    assert 'Reference:' in response.json()['detail']
    assert 'stage=extract_resume_data' in caplog.text
    assert 'error_type=RuntimeError' in caplog.text
    assert 'Provider unavailable' not in caplog.text
    assert db.query(Resume).count() == 0
    assert db.query(Subscription).first().scans_this_month == 0


def test_provider_auth_failure_is_identified_without_exposing_details(client, db, monkeypatch, caplog):
    import httpx
    from openai import AuthenticationError
    def denied(text):
        response = httpx.Response(401, request=httpx.Request('POST', 'https://example.com'))
        raise AuthenticationError('private-provider-response', response=response, body=None)
    monkeypatch.setattr('backend.app.utils.ats_engine.extract_resume_data', denied)
    input_pdf = generate_optimized_resume('input.pdf', resume_data=original()).getvalue()
    response = client.post('/api/v1/resume/upload', headers=headers(),
        files={'file': ('input.pdf', input_pdf, 'application/pdf')},
        data={'job_title': 'Engineer', 'job_description': 'Python experience required for this engineering role.'})
    assert response.status_code == 503
    assert 'credentials or permissions' in response.json()['detail']
    assert 'provider_status=401' in caplog.text
    assert 'private-provider-response' not in response.text + caplog.text
    assert db.query(Resume).count() == 0
    assert db.query(Subscription).first().scans_this_month == 0


def test_profile_read_does_not_fabricate_credentials(client,db):
    assert client.get('/api/v1/profile/me',headers=headers()).json()['certified_skills']==[]
    assert db.query(User).filter_by(id=2).first().languages==[]


def test_referral_cannot_convert_publicly_or_attribute_other_user(client,db):
    assert client.post('/api/v1/referral/convert?referred_user_id=2&plan=pro').status_code==404
    assert client.post('/api/v1/referral/track?ref_code=owner1&referred_email=other@example.com',headers=headers()).status_code==403
    assert client.post('/api/v1/referral/track?ref_code=arbitrary1&referred_email=candidate@example.com',headers=headers()).status_code==400


def test_score_is_repeatable_grounded_and_not_substring_match():
    jd='Python and Terraform required; Git preferred.'
    req=[{'phrase':'Python'},{'phrase':'Terraform'},{'phrase':'Git','priority':'preferred'},{'phrase':'Unlisted skill'}]
    first=score_requirements('Python developer. GitHub projects.',jd,req)
    assert first==score_requirements('Python developer. GitHub projects.',jd,req)
    assert first['overall_score']==40 and first['keywords_matched']==1
    assert first['total_keywords']==3


def original():
    return {'name':'Candidate','contact':'candidate@example.com','summary':'Python developer','experience':[{'role':'Engineer','company':'Example','dates':'2025','bullets':['Built Python tools.']}], 'projects':[{'title':'Project','bullets':['Built a tool.']}],'skills':{'Tools':['Python']},'education':[{'degree':'BSc','institution':'School','year':'2025'}],'certifications':['AWS certification - In progress']}


def test_rewrite_rejects_missing_history_new_skills_and_metrics():
    base=original(); rewrite={'optimized_summary':'Python developer','optimized_experience':base['experience'],'optimized_skills':base['skills']}
    source='Candidate candidate@example.com Engineer Example 2025 Python developer'
    assert merge_rewrite(base,rewrite,source)['education']==base['education']
    with pytest.raises(ValueError):merge_rewrite(base,{**rewrite,'optimized_experience':[]},source)
    with pytest.raises(ValueError):merge_rewrite(base,{**rewrite,'optimized_skills':{'Tools':['Terraform']}},source)
    with pytest.raises(ValueError):merge_rewrite(base,{**rewrite,'optimized_summary':'Improved output by 90%'},source)


def test_pdf_preserves_sections_and_omits_optimizer_notes():
    pdf=generate_optimized_resume('input.pdf',resume_data=original(),improvements=[{'bullet_point':'Added Terraform'}]).getvalue()
    text=extract_text(pdf,'output.pdf')
    assert all(word in text for word in ['PROJECTS','EDUCATION','CERTIFICATIONS','In progress'])
    assert 'Added Terraform' not in text and 'Terraform' not in text


def test_docx_table_text_is_retained():
    from docx import Document
    doc=Document();doc.add_paragraph('Candidate');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='School';table.cell(0,1).text='Degree'
    out=io.BytesIO();doc.save(out)
    assert 'School Degree' in extract_text(out.getvalue(),'resume.docx')


def test_scraper_keeps_full_text_and_structure():
    from backend.app.utils.global_scraper import _clean, _work_type
    content='<h2>Requirements</h2><ul><li>Python</li><li>Terraform</li></ul><p>'+('description '*600)+'</p>'
    clean=_clean(content)
    assert len(clean)>4000 and '\n' in clean and '- Python' in clean
    assert _work_type('','','')=='unknown'


def test_signed_payment_conversion_is_idempotent(db):
    from backend.app.routes.referral import convert_referral
    ref=Referral(referrer_id=1,referred_user_id=2,referred_email='candidate@example.com',status='pending')
    db.add(ref);db.commit()
    assert convert_referral(2,'pro',db)['reward']==5
    db.commit()
    assert convert_referral(2,'pro',db)['message']=='No pending referral found for this user'
    assert db.query(Referral).count()==1
